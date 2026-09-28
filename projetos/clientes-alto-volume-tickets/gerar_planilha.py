"""Gera a planilha "Analise Clientes - Volumetria de Tickets" a partir dos brutos em dados/. Regras e pesos: ver analise-tickets.ipynb."""
import warnings
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from definicoes_planilha import CATEGORIAS, ERRO_PRODUTO, montar_notas

warnings.filterwarnings("ignore")

PASTA = Path(__file__).parent
SAIDA = PASTA.parent.parent / "retencao" / "usuarios-por-produto-2026-09-24" / "Analise Clientes - Volumetria de Tickets.xlsx"

URGENCIA = {"Crítica": ("C0392B", "FFFFFF"), "Alta": ("E67E22", "FFFFFF"), "Média": ("F7DC6F", "000000"), "Baixa": ("D5F5E3", "000000")}
PARAMETROS = {
    "PesoValor": ("Peso do valor pago", 0.40, "0%"),
    "PesoErros": ("Peso dos erros de produto", 0.40, "0%"),
    "PesoVolume": ("Peso do volume total de tickets", 0.20, "0%"),
    "CorteCritica": ("Urgência Crítica a partir do score", 85, "0"),
    "CorteAlta": ("Urgência Alta a partir do score", 70, "0"),
    "CorteMedia": ("Urgência Média a partir do score", 45, "0"),
    "PercentilErroAlta": ("Cliente entre os X% com mais erros de produto é no mínimo Alta (percentil)", 0.90, "0%"),
    "DestaqueMinTickets": ("Categoria em destaque: mínimo de tickets do cliente no tipo", 5, "0"),
    "DestaqueFator": ("Categoria em destaque: quantas vezes acima da média da base", 1.5, '0.0"x"'),
}
ULTIMA_DATA_BASE = pd.Timestamp("2026-09-22")

FONTE = "Arial"
COR_TITULO = "0E4D4A"
F_NORMAL = Font(name=FONTE, size=10)
F_CAB = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
F_TITULO = Font(name=FONTE, size=14, bold=True, color=COR_TITULO)
F_NOTA = Font(name=FONTE, size=10, italic=True, color="555555")
F_SECAO = Font(name=FONTE, size=11, bold=True, color=COR_TITULO)
FMT_REAIS = '"R$" #,##0.00'
AMARELO = PatternFill("solid", fgColor="FFF2CC", bgColor="FFF2CC")


def fill(cor):
    # bgColor também: formatação condicional do Excel ignora fundo sólido definido só em fgColor
    return PatternFill("solid", fgColor=cor, bgColor=cor)


def categoria_analise(cat, causa, nec):
    if cat in ("", "Pesquisa Satisfação"):
        return "Fora da análise"
    if cat in ("Bug", "Solução de contorno"):
        return ERRO_PRODUTO
    if cat == "Problema":
        if causa.startswith("Bug no Produto"):
            return ERRO_PRODUTO
        return {"Configuração": "Problema de configuração", "Erro operacional": "Erro de uso"}.get(causa, "Externo ou sem causa definida")
    if cat == "Adequação":
        return "Adequação à legislação" if nec == "Legislação" else "Melhoria no produto"
    return {"Dúvida": "Dúvida", "Solicitação de serviço": "Solicitação de serviço"}.get(cat, "Fora da análise")


def percentil(s):
    return s.apply(lambda x: (s < x).sum()) / (len(s) - 1)


def cabecalho(ws, linha, titulos, cor=COR_TITULO, larguras=None):
    for i, t in enumerate(titulos, 1):
        c = ws.cell(linha, i, t)
        c.font, c.fill = F_CAB, fill(cor)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[linha].height = 45
    for i, w in enumerate(larguras or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def escrever(ws, linha, valores, formatos=None):
    for i, v in enumerate(valores, 1):
        c = ws.cell(linha, i, None if (v is None or (not isinstance(v, str) and pd.isna(v))) else v)
        c.font = F_NORMAL
        if formatos and formatos.get(i):
            c.number_format = formatos[i]


def regras_urgencia(ws, rng):
    for nome, (bg, fg) in URGENCIA.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{nome}"'], fill=fill(bg), font=Font(color=fg, bold=True)))


def regras_categoria(ws, rng):
    for nome, (forte, clara, _) in CATEGORIAS.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{nome}"'], fill=fill(clara), font=Font(color=forte, bold=True)))


def regra_sim(ws, col_sim, rng_destaque, primeira):
    ws.conditional_formatting.add(f"{col_sim}{primeira}:{col_sim}{ws.max_row}", CellIsRule(operator="equal", formula=['"Sim"'], fill=fill("1E8449"), font=Font(color="FFFFFF", bold=True)))
    ws.conditional_formatting.add(rng_destaque, FormulaRule(formula=[f'${col_sim}{primeira}="Sim"'], fill=AMARELO, font=Font(bold=True)))


# ---------- dados ----------
tickets = pd.read_excel(PASTA / "dados" / "Analise de tickets abertos.xlsx")
contratos = pd.read_excel(PASTA / "dados" / "Lista de Clientes.xlsx")
contratos = contratos[contratos["CONTRATO"].notna()].copy()
contratos["CODI_GRUPO"] = contratos["CODI_GRUPO"].astype("Int64")

for col in ("Categoria", "Causa", "Necessidade"):
    tickets[col] = tickets[col].fillna("")
tickets["cat_analise"] = [categoria_analise(a, b, c) for a, b, c in zip(tickets["Categoria"], tickets["Causa"], tickets["Necessidade"])]
tickets["entra"] = (tickets["Tipo"] == "Público") & (tickets["Status"] != "Cancelado") & (tickets["Categoria"] != "Pesquisa Satisfação")
tickets["em_aberto"] = ~tickets["Status"].isin(["Fechado", "Resolvido", "Cancelado"])
tickets["modulo"] = tickets["4° Nível Serviço"].fillna(tickets["3° Nível Serviço"]).fillna("não informado")
tickets = tickets.sort_values(["Abertura", "Ticket"], ascending=[False, False]).reset_index(drop=True)
tk = tickets[tickets["entra"]]

por_cliente = tk.pivot_table(index="ID Grupo Econômico", columns="cat_analise", values="Ticket", aggfunc="count", fill_value=0).reindex(columns=list(CATEGORIAS), fill_value=0)
por_cliente["total"] = por_cliente.sum(axis=1)
valor = contratos[contratos["SITUACAO_CONTRATO"] == "ATIVO"].groupby("CODI_GRUPO")["VALOR_CONTRATO"].sum()
por_cliente["valor"] = por_cliente.index.map(valor).fillna(0).astype(float)
pct = {c: percentil(por_cliente[c]) for c in ("valor", ERRO_PRODUTO, "total")}
por_cliente["score"] = 100 * (0.4 * pct["valor"] + 0.4 * pct[ERRO_PRODUTO] + 0.2 * pct["total"])
por_cliente = por_cliente.sort_values(["score", "valor"], ascending=False)

melhor = contratos.assign(ativo=contratos["SITUACAO_CONTRATO"] == "ATIVO").sort_values(["ativo", "VALOR_CONTRATO"], ascending=False).drop_duplicates("CODI_GRUPO").set_index("CODI_GRUPO")
nome_ticket = tk.groupby("ID Grupo Econômico")["Grupo Econônomico"].agg(lambda s: s.value_counts().index[0])
curva_ticket = tk.groupby("ID Grupo Econômico")["Curva"].agg(lambda s: s.value_counts().index[0])


def info(gid):
    if gid in melhor.index:
        m = melhor.loc[gid]
        return m["NOME_GRUPO"], m["CURVA"] if pd.notna(m["CURVA"]) else curva_ticket[gid], m["AGENTE_RELACIONAMENTO"], m["GERENTE_CONTA"]
    return nome_ticket[gid], curva_ticket[gid], None, None


wb = Workbook()
wb.calculation.fullCalcOnLoad = True
ws_score = wb.active
ws_score.title = "Score dos clientes"
ws_fila = wb.create_sheet("Fila de atendimento")
ws_cat = {c: wb.create_sheet(c) for c in CATEGORIAS}
ws_base = wb.create_sheet("Base de tickets")
ws_contr = wb.create_sheet("Contratos")
ws_par = wb.create_sheet("Parâmetros")

# ---------- Contratos ----------
cols_contr = ["CODI_GRUPO", "NOME_GRUPO", "COD_CLIENTE", "CLIENTE", "PRODUTO", "VALOR_CONTRATO", "PERIODICIDADE", "SITUACAO_CONTRATO", "DATA_CONTRATO", "CONTRATO", "CURVA", "GERENTE_CONTA", "AGENTE_RELACIONAMENTO", "CIDADE", "ESTADO", "SEGMENTO_PRINCIPAL"]
cabecalho(ws_contr, 1, cols_contr, "5D6D7E", [13, 40, 12, 40, 24, 15, 9, 13, 12, 12, 7, 30, 30, 20, 7, 40])
for r, row in enumerate(contratos[cols_contr].itertuples(index=False), 2):
    v = list(row)
    v[0] = int(v[0]) if pd.notna(v[0]) else None
    v[2] = str(v[2]) if pd.notna(v[2]) else None
    v[8] = v[8].to_pydatetime() if pd.notna(v[8]) else None
    v[9] = int(v[9]) if pd.notna(v[9]) else None
    escrever(ws_contr, r, v, {6: FMT_REAIS, 9: "dd/mm/yyyy"})
UC = ws_contr.max_row
ws_contr.freeze_panes = "C2"
ws_contr.auto_filter.ref = f"A1:P{UC}"
ws_contr.sheet_properties.tabColor = "5D6D7E"
C_GRUPO, C_PROD, C_VALOR, C_SIT = (f"Contratos!${c}$2:${c}${UC}" for c in "AEFH")

# ---------- Parâmetros ----------
ws_par.column_dimensions["A"].width = 64
ws_par.column_dimensions["B"].width = 22
ws_par.column_dimensions["C"].width = 22
ws_par.column_dimensions["D"].width = 26
ws_par.column_dimensions["E"].width = 30
ws_par["A1"] = "Como usar esta planilha"
ws_par["A1"].font = F_TITULO
instrucoes = [
    "1. Na aba \"Score dos clientes\", escolha \"Sim\" na coluna \"Prioridade CS\" para o cliente que o CS decidiu priorizar depois da análise humana.",
    "2. A aba \"Fila de atendimento\" se reordena sozinha: primeiro os clientes marcados com \"Sim\", depois do maior para o menor score.",
    "3. Cada categoria de ticket tem uma aba com cor própria: clientes que mais abrem aquele tipo de ticket e os módulos envolvidos.",
    "4. Score (0 a 100): posição do cliente em valor pago, erros de produto e volume de tickets, comparado aos demais, ponderada pelos pesos abaixo.",
    "   Categoria predominante: o tipo de ticket que o cliente mais abre (na maioria é Dúvida, que é metade de toda a base).",
    "   Categoria em destaque: o tipo em que o cliente mais foge da média da base (ex.: 3,4x mais erro de produto que a média). \"Perfil na média\" quando nenhum tipo se destaca.",
    "5. Base: tickets do AgriManager de 02/01/2025 a 22/09/2026 abertos pelo cliente (sem internos, cancelados e pesquisas de satisfação).",
    "6. Valor mensal: soma dos contratos ativos do grupo econômico, em todos os produtos. Cruzamento pelo código do grupo econômico.",
    "7. Células amarelas podem ser alteradas; score, urgência e fila se recalculam.",
]
for i, t in enumerate(instrucoes, 2):
    ws_par.cell(i, 1, t).font = F_NORMAL

lin = len(instrucoes) + 3
ws_par.cell(lin, 1, "Pesos e cortes").font = F_SECAO
lin += 1
for nome, (rotulo, valor_padrao, fmt) in PARAMETROS.items():
    ws_par.cell(lin, 1, rotulo).font = F_NORMAL
    c = ws_par.cell(lin, 2, valor_padrao)
    c.font, c.fill, c.number_format = Font(name=FONTE, size=10, color="0000FF"), AMARELO, fmt
    wb.defined_names[nome] = DefinedName(nome, attr_text=f"'Parâmetros'!$B${lin}")
    lin += 1
ws_par.cell(lin, 1, "Soma dos pesos (precisa dar 100%)").font = F_NORMAL
c = ws_par.cell(lin, 2, "=PesoValor+PesoErros+PesoVolume")
c.font, c.number_format = F_NORMAL, "0%"
lin += 1
ws_par.cell(lin, 1, "Última data da base de tickets").font = F_NORMAL
c = ws_par.cell(lin, 2, ULTIMA_DATA_BASE.to_pydatetime())
c.font, c.fill, c.number_format = Font(name=FONTE, size=10, color="0000FF"), AMARELO, "dd/mm/yyyy"
lin_data = lin
lin += 1
ws_par.cell(lin, 1, "Início da janela \"últimos 6 meses\"").font = F_NORMAL
c = ws_par.cell(lin, 2, f"=EDATE(B{lin_data},-6)")
c.font, c.number_format = F_NORMAL, "dd/mm/yyyy"
wb.defined_names["InicioJanela"] = DefinedName("InicioJanela", attr_text=f"'Parâmetros'!$B${lin}")

lin += 2
ws_par.cell(lin, 1, "Categorias de análise (cada uma tem sua cor e sua aba)").font = F_SECAO
lin += 1
cabecalho(ws_par, lin, ["Categoria", "Cor", "O que entra"])
for nome, (forte, clara, desc) in CATEGORIAS.items():
    lin += 1
    ws_par.cell(lin, 1, nome).font = Font(name=FONTE, size=10, bold=True, color=forte)
    ws_par.cell(lin, 2).fill = fill(forte)
    ws_par.cell(lin, 3, desc).font = F_NORMAL

lin += 2
ws_par.cell(lin, 1, "Regras de classificação (Categoria + Causa + Necessidade do ticket → categoria de análise)").font = F_SECAO
lin += 1
cabecalho(ws_par, lin, ["Chave (automática)", "Categoria", "Causa", "Necessidade", "Categoria de análise"])
combos = tickets[["Categoria", "Causa", "Necessidade", "cat_analise"]].drop_duplicates().sort_values(["cat_analise", "Categoria", "Causa", "Necessidade"])
r0 = lin + 1
for r, (cat, causa, nec, ca) in enumerate(combos.itertuples(index=False), r0):
    escrever(ws_par, r, [f'=B{r}&"|"&C{r}&"|"&D{r}', cat or None, causa or None, nec or None, ca])
    ws_par.cell(r, 5).fill = AMARELO
r1 = r0 + len(combos) - 1
REGRA_CHAVE, REGRA_CAT = f"'Parâmetros'!$A${r0}:$A${r1}", f"'Parâmetros'!$E${r0}:$E${r1}"
ws_par.sheet_properties.tabColor = "F1C40F"

# ---------- Base de tickets ----------
cols_base = ["Ticket", "Abertura", "ID grupo econômico", "Grupo econômico (nome no ticket)", "Tipo", "Status", "Prioridade", "Categoria", "Causa", "Necessidade", "Categoria de análise", "Entra na análise", "Em aberto", "Módulo (4º nível)", "Área (3º nível)", "Squad", "Assunto", "Avaliação"]
cabecalho(ws_base, 1, cols_base, "5D6D7E", [9, 11, 12, 40, 9, 16, 10, 20, 26, 20, 26, 10, 9, 45, 30, 34, 60, 10])
base_df = tickets[["Ticket", "Abertura", "ID Grupo Econômico", "Grupo Econônomico", "Tipo", "Status", "Prioridade", "Categoria", "Causa", "Necessidade", "4° Nível Serviço", "3° Nível Serviço", "Squad", "Assunto", "Avaliação"]]
for r, v in enumerate(base_df.itertuples(index=False, name=None), 2):
    v = list(v)
    escrever(ws_base, r, [
        int(v[0]), v[1].to_pydatetime(), int(v[2]), v[3], v[4], v[5], v[6],
        v[7] or None, v[8] or None, v[9] or None,
        f'=IFERROR(INDEX({REGRA_CAT},MATCH(H{r}&"|"&I{r}&"|"&J{r},{REGRA_CHAVE},0)),"Sem regra")',
        f'=IF(AND(E{r}="Público",F{r}<>"Cancelado",H{r}<>"Pesquisa Satisfação"),"Sim","Não")',
        f'=IF(OR(F{r}="Fechado",F{r}="Resolvido",F{r}="Cancelado"),"Não","Sim")',
        v[10], v[11], v[12], v[13], v[14],
    ], {2: "dd/mm/yyyy"})
UB = ws_base.max_row
ws_base.freeze_panes = "B2"
ws_base.auto_filter.ref = f"A1:R{UB}"
ws_base.sheet_properties.tabColor = "5D6D7E"
B = {k: f"'Base de tickets'!${c}$2:${c}${UB}" for k, c in {"abertura": "B", "grupo": "C", "cat": "K", "entra": "L", "aberto": "M"}.items()}

# ---------- Score dos clientes ----------
marcados_cs = set()
if SAIDA.exists():
    linhas_ant = load_workbook(SAIDA, read_only=True)["Score dos clientes"].iter_rows(values_only=True)
    cab_ant = next(linhas_ant)
    i_prio, i_id = cab_ant.index("Prioridade CS"), cab_ant.index("ID grupo econômico")
    marcados_cs = {lin[i_id] for lin in linhas_ant if str(lin[i_prio]).strip().lower() == "sim"}
n = len(por_cliente)
US = n + 1
FATORES = [f"Fator {c} (auxiliar)" for c in CATEGORIAS]
cols_score = (
    ["Prioridade CS", "Cliente (grupo econômico)", "Valor mensal", "Score", "Urgência", "Categoria em destaque",
     "Vezes acima da média", "Categoria predominante", "Total de tickets", "% erro de produto"]
    + list(CATEGORIAS)
    + ["Em aberto hoje", "Tickets nos últimos 6 meses", "Situação AgriManager", "Valor mensal AgriManager", "Curva",
       "CS responsável", "Gerente de conta", "ID grupo econômico", "Ordem (auxiliar)"]
    + FATORES
)
L = {nome: get_column_letter(i) for i, nome in enumerate(cols_score, 1)}
CAT1, CATN = L[ERRO_PRODUTO], L[list(CATEGORIAS)[-1]]
FAT1, FATN = L[FATORES[0]], L[FATORES[-1]]
larg = {"Prioridade CS": 11, "Cliente (grupo econômico)": 52, "Valor mensal": 14, "Score": 9, "Urgência": 11,
        "Categoria em destaque": 28, "Vezes acima da média": 11, "Categoria predominante": 24,
        "Situação AgriManager": 22, "Valor mensal AgriManager": 14, "Curva": 7, "CS responsável": 38, "Gerente de conta": 36}
cabecalho(ws_score, 1, cols_score, larguras=[larg.get(c, 12) for c in cols_score])
for nome, (forte, _, _) in CATEGORIAS.items():
    ws_score[f"{L[nome]}1"].fill = fill(forte)
col = lambda nome: f"${L[nome]}$2:${L[nome]}${US}"
fmt_score = {i: f for i, f in enumerate([None] * len(cols_score), 1)}
for nome, f in {"Valor mensal": FMT_REAIS, "Score": "0.0", "Vezes acima da média": '0.0"x"', "% erro de produto": "0%",
                "Valor mensal AgriManager": FMT_REAIS, "Ordem (auxiliar)": "0.0000", **{c: "0.00" for c in FATORES}}.items():
    fmt_score[cols_score.index(nome) + 1] = f
for r, gid in enumerate(por_cliente.index, 2):
    nome, curva, cs, gerente = info(gid)
    x, tot, err, sc_ = f"${L['ID grupo econômico']}{r}", f"{L['Total de tickets']}{r}", f"{L[ERRO_PRODUTO]}{r}", f"{L['Score']}{r}"
    fatores = f"{FAT1}{r}:{FATN}{r}"
    f = {
        "Prioridade CS": "Sim" if int(gid) in marcados_cs else "Não",
        "Cliente (grupo econômico)": nome,
        "Valor mensal": f'=SUMIFS({C_VALOR},{C_GRUPO},{x},{C_SIT},"ATIVO")',
        "Score": f"=100*(PesoValor*PERCENTRANK({col('Valor mensal')},{L['Valor mensal']}{r},6)+PesoErros*PERCENTRANK({col(ERRO_PRODUTO)},{err},6)+PesoVolume*PERCENTRANK({col('Total de tickets')},{tot},6))",
        "Urgência": f'=IF({sc_}>=CorteCritica,"Crítica",IF(OR({sc_}>=CorteAlta,PERCENTRANK({col(ERRO_PRODUTO)},{err},6)>=PercentilErroAlta),"Alta",IF({sc_}>=CorteMedia,"Média","Baixa")))',
        "Categoria em destaque": f'=IF(MAX({fatores})<DestaqueFator,"Perfil na média",INDEX(${CAT1}$1:${CATN}$1,MATCH(MAX({fatores}),{fatores},0)))',
        "Vezes acima da média": f'=IF(MAX({fatores})<DestaqueFator,"",MAX({fatores}))',
        "Categoria predominante": f"=INDEX(${CAT1}$1:${CATN}$1,MATCH(MAX({CAT1}{r}:{CATN}{r}),{CAT1}{r}:{CATN}{r},0))",
        "Total de tickets": f"=SUM({CAT1}{r}:{CATN}{r})",
        "% erro de produto": f"=IF({tot}=0,0,{err}/{tot})",
        "Em aberto hoje": f'=COUNTIFS({B["grupo"]},{x},{B["entra"]},"Sim",{B["aberto"]},"Sim")',
        "Tickets nos últimos 6 meses": f'=COUNTIFS({B["grupo"]},{x},{B["entra"]},"Sim",{B["abertura"]},">="&InicioJanela)',
        "Situação AgriManager": f'=IF(COUNTIFS({C_GRUPO},{x},{C_PROD},"AGRIMANAGER*",{C_SIT},"ATIVO")>0,"Ativo",IF(COUNTIFS({C_GRUPO},{x},{C_PROD},"AGRIMANAGER*",{C_SIT},"SUSPENSO")>0,"Suspenso",IF(COUNTIFS({C_GRUPO},{x},{C_PROD},"AGRIMANAGER*")>0,"Cancelado",IF(COUNTIF({C_GRUPO},{x})>0,"Sem contrato AgriManager","Fora da lista de clientes"))))',
        "Valor mensal AgriManager": f'=SUMIFS({C_VALOR},{C_GRUPO},{x},{C_PROD},"AGRIMANAGER*",{C_SIT},"ATIVO")',
        "Curva": curva, "CS responsável": cs, "Gerente de conta": gerente, "ID grupo econômico": int(gid),
        "Ordem (auxiliar)": f'=IF($A{r}="Sim",1000,0)+{sc_}+ROW()/10000000',
    }
    for c in CATEGORIAS:
        f[c] = f'=COUNTIFS({B["grupo"]},{x},{B["cat"]},{L[c]}$1,{B["entra"]},"Sim")'
    for c, fat in zip(CATEGORIAS, FATORES):
        f[fat] = f"=IF({L[c]}{r}<DestaqueMinTickets,0,({L[c]}{r}/{tot})/(SUM({col(c)})/SUM({col('Total de tickets')})))"
    escrever(ws_score, r, [f[c] for c in cols_score], fmt_score)
dv = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True, showDropDown=False)
dv.error, dv.errorTitle = "Escolha Sim ou Não", "Prioridade CS"
ws_score.add_data_validation(dv)
dv.add(f"A2:A{US}")
ws_score.freeze_panes = "D2"
ws_score.auto_filter.ref = f"A1:{L['Ordem (auxiliar)']}{US}"
for c in ["Ordem (auxiliar)"] + FATORES:
    ws_score.column_dimensions[L[c]].hidden = True
ws_score.sheet_properties.tabColor = COR_TITULO
regra_sim(ws_score, "A", f"B2:C{US}", 2)
regras_urgencia(ws_score, f"{L['Urgência']}2:{L['Urgência']}{US}")
regras_categoria(ws_score, f"{L['Categoria em destaque']}2:{L['Categoria em destaque']}{US}")
regras_categoria(ws_score, f"{L['Categoria predominante']}2:{L['Categoria predominante']}{US}")
ws_score.conditional_formatting.add(f"{L['Score']}2:{L['Score']}{US}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color="5DADE2"))
for nome, (forte, _, _) in CATEGORIAS.items():
    ws_score.conditional_formatting.add(f"{L[nome]}2:{L[nome]}{US}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="max", end_color=forte))

# ---------- Fila de atendimento ----------
ws_fila["A1"] = "Fila de atendimento"
ws_fila["A1"].font = F_TITULO
ws_fila["A2"] = "Ordem automática: primeiro os clientes marcados com \"Sim\" em Prioridade CS (aba Score dos clientes), depois do maior para o menor score. Não editar esta aba."
ws_fila["A2"].font = F_NOTA
cols_fila = ["Posição", "Prioridade CS", "Cliente (grupo econômico)", "Valor mensal", "Score", "Urgência", "Categoria em destaque",
             "Vezes acima da média", "Categoria predominante", ERRO_PRODUTO, "Total de tickets", "Em aberto hoje", "CS responsável", "Linha (auxiliar)"]
LF = {nome: get_column_letter(i) for i, nome in enumerate(cols_fila, 1)}
cabecalho(ws_fila, 4, cols_fila, larguras=[8, 11, 52, 14, 9, 11, 28, 11, 24, 10, 10, 10, 38, 8])
S = lambda nome: f"'Score dos clientes'!${L[nome]}$2:${L[nome]}${US}"
for k in range(1, n + 1):
    r = k + 4
    lin_aux = f"${LF['Linha (auxiliar)']}{r}"
    v = [k] + [f"=INDEX({S(c)},{lin_aux})" for c in cols_fila[1:-2]] + [
        f'=IF(INDEX({S("CS responsável")},{lin_aux})=0,"",INDEX({S("CS responsável")},{lin_aux}))',
        f"=MATCH(LARGE({S('Ordem (auxiliar)')},A{r}),{S('Ordem (auxiliar)')},0)",
    ]
    escrever(ws_fila, r, v, {4: FMT_REAIS, 5: "0.0", 8: '0.0"x"'})
UF = n + 4
ws_fila.freeze_panes = "D5"
ws_fila.column_dimensions[LF["Linha (auxiliar)"]].hidden = True
ws_fila.sheet_properties.tabColor = "1E8449"
regra_sim(ws_fila, "B", f"C5:D{UF}", 5)
regras_urgencia(ws_fila, f"F5:F{UF}")
regras_categoria(ws_fila, f"G5:G{UF}")
regras_categoria(ws_fila, f"I5:I{UF}")
ws_fila.conditional_formatting.add(f"E5:E{UF}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color="5DADE2"))

# ---------- uma aba por categoria ----------
for nome, (forte, clara, desc) in CATEGORIAS.items():
    ws = ws_cat[nome]
    ws["A1"] = nome
    ws["A1"].font = Font(name=FONTE, size=14, bold=True, color=forte)
    ws["A2"] = f"O que entra: {desc} Ordenado pelo número de tickets desta categoria."
    ws["A2"].font = F_NOTA
    cabecalho(ws, 4, ["Posição", "Cliente (grupo econômico)", "Valor mensal", "Tickets nesta categoria", "Total de tickets do cliente", "% desta categoria no cliente", "Em aberto nesta categoria", "Urgência do cliente", "Prioridade CS", "Principais módulos", "Ticket mais recente", "ID grupo econômico"], forte, [8, 52, 14, 11, 11, 11, 11, 11, 11, 100, 90, 12])
    sub = tk[tk["cat_analise"] == nome]
    ranking = por_cliente[por_cliente[nome] > 0].sort_values([nome, "valor"], ascending=False)
    for k, gid in enumerate(ranking.index, 1):
        r = k + 4
        st = sub[sub["ID Grupo Econômico"] == gid]
        mods = "; ".join(f"{m} ({q})" for m, q in st["modulo"].value_counts().head(3).items())
        ult = st.iloc[0]
        idx = f"MATCH($L{r},{S('ID grupo econômico')},0)"
        escrever(ws, r, [
            k, info(gid)[0], f"=INDEX({S('Valor mensal')},{idx})",
            f'=COUNTIFS({B["grupo"]},$L{r},{B["cat"]},$A$1,{B["entra"]},"Sim")', f"=INDEX({S('Total de tickets')},{idx})",
            f"=IF(E{r}=0,0,D{r}/E{r})", f'=COUNTIFS({B["grupo"]},$L{r},{B["cat"]},$A$1,{B["entra"]},"Sim",{B["aberto"]},"Sim")',
            f"=INDEX({S('Urgência')},{idx})", f"=INDEX({S('Prioridade CS')},{idx})", mods,
            f"{ult['Abertura']:%d/%m/%Y}: {ult['Assunto']}", int(gid),
        ], {3: FMT_REAIS, 6: "0%"})
    u = len(ranking) + 4
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:L{u}"
    ws.sheet_properties.tabColor = forte
    ws.conditional_formatting.add(f"D5:D{u}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="max", end_color=forte))
    regras_urgencia(ws, f"H5:H{u}")
    regra_sim(ws, "I", f"B5:C{u}", 5)

# ---------- notas explicando cada coluna ----------
participacao = tk["cat_analise"].value_counts(normalize=True).to_dict()
for nota in montar_notas(participacao):
    ws = wb[nota["aba"]]
    linhas = [nota["linha"]] if nota["linha"] else range(1, ws.max_row + 1)
    achada = next((c for r in linhas for c in ws[r] if c.value == nota["procurar"]), None)
    if achada is None:
        raise ValueError(f"nota sem coluna: {nota['aba']} / {nota['procurar']}")
    texto = nota["nota"]
    altura = sum(max(1, -(-len(p) // 55)) for p in texto.split("\n")) * 13 + 14
    ws.cell(achada.row, achada.column + nota["deslocamento"]).comment = Comment(texto, "Análise de tickets", height=altura, width=330)

wb.save(SAIDA)
print(f"ok: {SAIDA.name} | clientes={n} | tickets base={UB - 1} | contratos={UC - 1} | Prioridade CS mantidos={len(marcados_cs)}")
