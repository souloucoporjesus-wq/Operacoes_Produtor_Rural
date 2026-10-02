"""Modelo de dados dos projetos de implantação hunter (myFarm e AgriManager).

Lê a exportação do CES (planilha com as abas Projetos, Hr Adquirida, OS_EXECUTADA,
OS PLANEJADA, IMPLANTADOR e STATUS_ATUALIZACAO) e cruza as abas:

    Projetos.COD ── Hr Adquirida.PROJETO          (pacotes de horas do projeto)
    Projetos.COD ── OS_EXECUTADA.COD_PROJETO      (cada linha = um dia de agenda)
    Hr Adquirida.CODIGO ── OS_EXECUTADA.COD_HORA_ADQUIRIDA
    OS_EXECUTADA.OS ── OS PLANEJADA.OS            (dias planejados da OS)
    OS_EXECUTADA.COD_CONSULTOR ── IMPLANTADOR.COD_IMPLANTADOR

O notebook de análise e o gerador do painel usam este módulo, então os números batem.
"""
from dataclasses import dataclass
from pathlib import Path
import re

import pandas as pd

# ---------------------------------------------------------------- regras (combinadas em 02/10/2026)
TIPOS_HUNTER = {"IMPLANTACAO", "REIMPLANTACAO"}
DEPTOS_HUNTER = {"CES - IMPLANTACAO AGRIMANAGER", "CES - IMPLANTACAO MYFARM"}
STATUS_ABERTO = ["EM ANDAMENTO", "SUSPENSO", "NAO INICIADO", "PENDENCIA"]
STATUS_HORA_UTILIZADA = {"FECHADA", "VALIDADA", "ACERTADA", "CONCLUIDA"}
SISTEMAS = {"MYFARM": "myFarm", "SIAGRI AGRIMANAGER": "AgriManager"}

# limites da saúde do projeto (dias e percentuais); ajuste aqui se quiser outro corte
DIAS_SEM_AGENDA_ATENCAO = 15
DIAS_SEM_AGENDA_CRITICO = 30
CONSUMO_ATENCAO = 0.75   # consumiu 75% das horas sem ter chegado ao go live
CONSUMO_CRITICO = 0.90   # consumiu 90% das horas sem ter chegado ao go live

FASES = {
    1: "Iniciação",
    2: "Planejamento",
    3: "Parametrização",
    4: "Treinamento e simulação",
    5: "Go live",
    6: "Pós produção",
    7: "Encerramento",
}
FASE_GO_LIVE = 5

GRUPO_HORA = {
    "NORMAL": "Contratadas",
    "GERENCIAMENTO DE PROJETO": "Contratadas",
    "INSTALACAO BANCO E SISTEMA": "Contratadas",
    "SERVICOS DBA": "Contratadas",
    "ACOMPANHAMENTO": "Contratadas",
    "COMPLEMENTAR": "Contratadas",
    "EXCEDENTE": "Contratadas",
    "BONFICACAO COM PRODUTIVIDADE": "Bonificadas",
    "BONIFICACAO SEM PRODUTIVIDADE": "Bonificadas",
    "TRANSFERENCIA": "Transferidas",
    "CREDITO FINANCEIRO": "Crédito financeiro",
}


def fase_da_etapa(etapa):
    """Traduz a etapa da OS para a fase da jornada (0 = gestão do projeto, sem fase)."""
    if not isinstance(etapa, str):
        return 0
    e = etapa.upper()
    regras = [
        (7, r"ENCERRAMENTO|TERMO DE ENCERRAMENTO|GANHOS E RESULTADOS|HANDOFF"),
        (6, r"P[OÓ]S PRODU|ACOMPANHAMENTO P[OÓ]S|MONITORIA"),
        (5, r"GO LIVE|ENTRADA EM PRODU"),
        (4, r"SIMULA|TREINAMENTO|HOMOLOGA"),
        (3, r"PARAMETRIZA"),
        (2, r"PLANEJAMENTO|LRN|LEVANTAMENTO|INSTALA[CÇ][AÃ]O|ESCOPO|MODELAGEM"),
        (1, r"INICIA[CÇ][AÃ]O|PATROCINADOR|TAP|APRESENTA[CÇ][AÃ]O DE TIMES|KICK"),
        (0, r"STATUS REPORT|MONITORAMENTO|GERENCIA"),
    ]
    for fase, padrao in regras:
        if re.search(padrao, e):
            return fase
    return 0


# ---------------------------------------------------------------- leitura
def _ler_aba(caminho, aba):
    df = pd.read_excel(caminho, sheet_name=aba, skiprows=2)
    df.columns = [re.sub(r".*\[(.*)\]", r"\1", str(c)) for c in df.columns]
    return df


def ler_planilha(caminho):
    """Lê as abas da exportação do CES, com os nomes de coluna limpos."""
    bruto = {
        "projetos": _ler_aba(caminho, "Projetos"),
        "horas": _ler_aba(caminho, "Hr Adquirida"),
        "os": _ler_aba(caminho, "OS_EXECUTADA"),
        "planejada": _ler_aba(caminho, "OS PLANEJADA"),
        "implantador": _ler_aba(caminho, "IMPLANTADOR"),
    }
    status = pd.read_excel(caminho, sheet_name="STATUS_ATUALIZACAO", header=None)
    bruto["atualizado_em"] = pd.to_datetime(status.iloc[0, 1])
    return bruto


ABAS_CES = {"Projetos", "Hr Adquirida", "OS_EXECUTADA", "OS PLANEJADA", "STATUS_ATUALIZACAO"}


def eh_exportacao_ces(arquivo):
    """True se a planilha tem as abas da exportação do CES (os reports de projetos não têm)."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(arquivo, read_only=True)
        abas = set(wb.sheetnames)
        wb.close()
    except Exception:
        return False
    return ABAS_CES <= abas


def planilha_mais_recente(pasta):
    """A exportação do CES mais nova da pasta (ignora reports e arquivos temporários do Excel)."""
    xs = sorted((x for x in Path(pasta).glob("*.xlsx") if not x.name.startswith("~$")),
                key=lambda x: x.stat().st_mtime, reverse=True)
    for x in xs:
        if eh_exportacao_ces(x):
            return x
    raise FileNotFoundError(f"nenhuma exportação do CES (.xlsx com as abas {', '.join(sorted(ABAS_CES))}) em {pasta}")


# ---------------------------------------------------------------- modelo
@dataclass
class Modelo:
    ref: pd.Timestamp            # data de posição (dia da exportação)
    atualizado_em: pd.Timestamp
    hunter: pd.DataFrame         # todos os projetos hunter (qualquer status)
    abertos: pd.DataFrame        # hunter em aberto, com todas as métricas
    agendas: pd.DataFrame        # agendas dos projetos hunter (realizadas, sem fechar, próximas, vencidas)
    horas: pd.DataFrame          # pacotes de horas adquiridas dos projetos hunter
    referencia: pd.DataFrame     # duração e horas típicas de projetos hunter concluídos desde 2024
    agenda_geral: pd.DataFrame   # agendas de todos os projetos (hunter e farmer), pra agenda dos consultores
    fases_referencia: dict       # parte das horas por fase nos hunter concluídos desde 2024, por sistema


def _projetos_hunter(proj):
    p = (proj.sort_values("MANUTENCAO_TOTAL_DO_GRUPO", na_position="first")
             .drop_duplicates("COD", keep="last"))
    p = p[p.TIPO_PROJETO.isin(TIPOS_HUNTER) & p.DESCRICAO_DEPARTAMENTO.isin(DEPTOS_HUNTER)].copy()
    p["SISTEMA_NOME"] = p.SISTEMA.map(SISTEMAS).fillna(p.SISTEMA)
    # DT_FINAL é o fim da execução; o encerramento formal (entrega pro suporte) vem em DT_ENTREGA_SUPORTE
    p["DT_ENCERRAMENTO"] = p.DT_ENTREGA_SUPORTE.fillna(p.DT_FINAL)
    return p


# horários das OS: pares (início, fim) de cada período; 00:00 no CES quer dizer campo vazio
PERIODOS_EXECUTADA = [("PRIMEIRA_ENTRADA_DIURNO", "PRIMEIRA_SAIDA_DIURNO"),
                      ("SEGUNDA_ENTRADA_DIURNO", "SEGUNDA_SAIDA_DIURNO"),
                      ("ENTRADA_NOTURNO", "SAIDA_NOTURNO")]
PERIODOS_PLANEJADA = [("HORA_INICIO_MANHA", "HORA_FIM_MANHA"), ("HORA_INICIO_TARDE", "HORA_FIM_TARDE")]


def _hhmm(t):
    """'08:00' a partir do horário do CES; vazio ou meia-noite vira None."""
    if isinstance(t, str):
        try:
            t = pd.to_datetime(t).time()
        except (ValueError, TypeError):
            return None
    if not hasattr(t, "hour") or (t.hour == 0 and t.minute == 0):
        return None
    return f"{t.hour:02d}:{t.minute:02d}"


def _periodos(df, pares):
    """Lista de (início, fim) de cada linha, só dos períodos preenchidos."""
    out = [[] for _ in range(len(df))]
    for ini, fim in pares:
        for i, (a, b) in enumerate(zip(df[ini].map(_hhmm), df[fim].map(_hhmm))):
            if isinstance(a, str) and isinstance(b, str):
                out[i].append((a, b))
    return pd.Series(out, index=df.index)


def _agendas(os_exec, planejada, cods, ref):
    """Uma linha por dia de agenda, com a situação dela na data de posição."""
    os_exec = os_exec[os_exec.COD_PROJETO.isin(cods)].copy()
    os_exec["DIA"] = os_exec.DATA_OS
    st = os_exec.STATUS_OS.fillna("")

    realizada = st.isin(STATUS_HORA_UTILIZADA) & os_exec.DIA.notna()
    confirmada = (st == "CONFIRMADA") & os_exec.DIA.notna()
    os_exec.loc[realizada, "SITUACAO"] = "Realizada"
    os_exec.loc[confirmada & (os_exec.DIA <= ref), "SITUACAO"] = "Sem fechar"
    os_exec.loc[confirmada & (os_exec.DIA > ref), "SITUACAO"] = "Próxima"
    feitas = os_exec[os_exec.SITUACAO.notna()].copy()
    feitas["HORAS"] = feitas.HORAS_TRABALHADAS.fillna(0)
    feitas["PERIODOS"] = _periodos(feitas, PERIODOS_EXECUTADA)

    # OS ainda não executadas: os dias vêm da OS PLANEJADA (ou da data de início da OS)
    pend = os_exec[st.isin({"PLANEJADA", "CONFIRMADA"}) & os_exec.DIA.isna()]
    pend = pend.drop_duplicates("OS")
    dias = planejada[planejada.OS.isin(pend.OS)].dropna(subset=["DATA_PLANEJADA"]).copy()
    dias["PERIODOS"] = _periodos(dias, PERIODOS_PLANEJADA)
    # o mesmo dia de uma OS pode vir partido em mais de uma linha (ex.: 14h às 17h e 17h às 18h)
    dias = dias.groupby(["OS", "DATA_PLANEJADA"], as_index=False).agg(
        HORAS_PLANEJADAS_DIARIO=("HORAS_PLANEJADAS_DIARIO", "sum"),
        PERIODOS=("PERIODOS", lambda v: sorted(p for lista in v for p in lista)))
    pend = pend.merge(dias, on="OS", how="left")
    pend["DIA"] = pend.DATA_PLANEJADA.fillna(pend.DATA_INICIO)
    pend["HORAS"] = pend.HORAS_PLANEJADAS_DIARIO.fillna(0)
    pend["PERIODOS"] = pend.PERIODOS.map(lambda v: v if isinstance(v, list) else [])
    pend["SITUACAO"] = (pend.DIA >= ref).map({True: "Próxima", False: "Vencida"})
    pend = pend.drop_duplicates(["OS", "DIA"])

    cols = ["OS", "COD_PROJETO", "PROJETO", "DIA", "CONSULTOR", "COD_CONSULTOR", "HORAS", "EXECUCAO",
            "ETAPA", "STATUS_OS", "SITUACAO", "COD_HORA_ADQUIRIDA", "PERIODOS"]
    ag = pd.concat([feitas[cols], pend[cols]], ignore_index=True)
    ag["FASE"] = ag.ETAPA.map(fase_da_etapa)
    ag["INICIO"] = ag.PERIODOS.map(lambda v: min(p[0] for p in v) if v else None)
    ag["FIM"] = ag.PERIODOS.map(lambda v: max(p[1] for p in v) if v else None)
    ag["HORARIO"] = ag.PERIODOS.map(lambda v: " e ".join(f"{a} às {b}" for a, b in v))
    return ag.sort_values(["COD_PROJETO", "DIA"]).reset_index(drop=True)


def _referencia(hunter, agendas_todas):
    """Duração (abertura até o encerramento formal) e horas usadas dos hunter concluídos desde 2024."""
    c = hunter[(hunter.STATUS == "CONCLUIDO") & (hunter.DT_ENCERRAMENTO >= "2024-01-01")].copy()
    c["DIAS"] = (c.DT_ENCERRAMENTO - c.DT_INICIAL).dt.days
    usadas = (agendas_todas[agendas_todas.SITUACAO == "Realizada"]
              .groupby("COD_PROJETO").HORAS.sum())
    c["HORAS_USADAS"] = c.COD.map(usadas).fillna(0)
    # ritmo: dias da abertura até a primeira agenda e a maior pausa entre duas agendas que aconteceram
    feitas = agendas_todas[agendas_todas.SITUACAO.isin(["Realizada", "Sem fechar"]) & agendas_todas.COD_PROJETO.isin(c.COD)]
    dias = feitas.groupby("COD_PROJETO").DIA.apply(lambda s: sorted(pd.Timestamp(x) for x in s.dt.normalize().unique()))
    c["ATE_PRIMEIRA"] = [(dias[k][0] - ini).days if k in dias.index else None for k, ini in zip(c.COD, c.DT_INICIAL)]
    c["MAIOR_PAUSA"] = [max(((b - a).days for a, b in zip(dias[k][:-1], dias[k][1:])), default=0) if k in dias.index else None
                        for k in c.COD]
    return (c.groupby("SISTEMA_NOME")
             .agg(projetos=("COD", "size"),
                  dias_mediana=("DIAS", "median"),
                  dias_p75=("DIAS", lambda s: s.quantile(0.75)),
                  horas_mediana=("HORAS_USADAS", "median"),
                  previsto_mediana=("PREVISTO", "median"),
                  ate_primeira_mediana=("ATE_PRIMEIRA", "median"),
                  maior_pausa_mediana=("MAIOR_PAUSA", "median")))


def _fases_referencia(hunter, agendas_todas):
    """Parte das horas realizadas em cada fase, somando os hunter concluídos desde 2024, por sistema."""
    c = hunter[(hunter.STATUS == "CONCLUIDO") & (hunter.DT_ENCERRAMENTO >= "2024-01-01")]
    r = agendas_todas[(agendas_todas.SITUACAO == "Realizada") & agendas_todas.COD_PROJETO.isin(c.COD)]
    r = r.merge(c[["COD", "SISTEMA_NOME"]], left_on="COD_PROJETO", right_on="COD")
    t = r.pivot_table(index="SISTEMA_NOME", columns="FASE", values="HORAS", aggfunc="sum", fill_value=0)
    t = t.div(t.sum(axis=1), axis=0)
    return {s: {int(k): round(float(v), 4) for k, v in linha.items()} for s, linha in t.iterrows()}


def _h(v):
    """Horas no formato brasileiro: 8h, 12,5h."""
    return f"{v:.1f}".replace(".", ",").removesuffix(",0") + "h"


SAUDES = ["Em dia", "Atenção", "Encerrar", "Crítico"]   # em ordem de gravidade; "Suspenso" fica à parte


def _saude(r, ref_sis):
    """Classifica o projeto e devolve os motivos (texto curto pra mostrar no painel).

    Crítico  = parado antes do go live, horas no fim sem go live, ou horas estouradas.
    Encerrar = já passou do go live e ficou sem agenda: falta encerrar ou usar o saldo.
    Atenção  = sinais menores (agenda vencida, OS sem fechar, poucos dias parado, aberto além do normal).
    """
    motivos, nivel = [], 0
    if r.STATUS == "SUSPENSO":
        return "Suspenso", ["projeto suspenso"]
    sem_proxima = pd.isna(r.PROXIMA_AGENDA)
    antes_go_live = r.FASE_ATUAL < FASE_GO_LIVE
    parado = r.DIAS_SEM_AGENDA > DIAS_SEM_AGENDA_CRITICO and sem_proxima

    if r.SALDO < 0:
        nivel = 3; motivos.append(f"usou {_h(abs(r.SALDO))} além do adquirido")
    if parado and antes_go_live:
        nivel = 3; motivos.append(f"parado antes do go live: {r.DIAS_SEM_AGENDA:.0f} dias sem agenda e nada marcado")
    if antes_go_live and r.ADQUIRIDAS > 0 and r.PCT_CONSUMO >= CONSUMO_CRITICO:
        nivel = 3; motivos.append(f"usou {r.PCT_CONSUMO:.0%} das horas sem chegar ao go live")
    if parado and not antes_go_live:
        nivel = max(nivel, 2)
        motivos.append(f"go live feito e {r.DIAS_SEM_AGENDA:.0f} dias sem agenda: encerrar ou usar o saldo de {_h(max(r.SALDO, 0))}")
    if DIAS_SEM_AGENDA_ATENCAO < r.DIAS_SEM_AGENDA <= DIAS_SEM_AGENDA_CRITICO and sem_proxima:
        nivel = max(nivel, 1); motivos.append(f"{r.DIAS_SEM_AGENDA:.0f} dias sem agenda e nada marcado")
    if antes_go_live and r.ADQUIRIDAS > 0 and CONSUMO_ATENCAO <= r.PCT_CONSUMO < CONSUMO_CRITICO:
        nivel = max(nivel, 1); motivos.append(f"usou {r.PCT_CONSUMO:.0%} das horas sem chegar ao go live")
    if r.AGENDAS_VENCIDAS > 0:
        nivel = max(nivel, 1); motivos.append(f"{r.AGENDAS_VENCIDAS:.0f} agenda(s) planejada(s) passaram sem execução nem cancelamento")
    if r.HORAS_SEM_FECHAR > 0:
        nivel = max(nivel, 1); motivos.append(f"{_h(r.HORAS_SEM_FECHAR)} de agenda confirmada com OS sem fechar")
    p75 = ref_sis.get(r.SISTEMA_NOME)
    if p75 is not None and r.DIAS_EM_ABERTO > p75:
        nivel = max(nivel, 1); motivos.append(f"aberto há {r.DIAS_EM_ABERTO:.0f} dias; 75% dos concluídos fecham em até {p75:.0f}")
    return SAUDES[nivel], motivos


# ordem de ataque: primeiro quem corre risco de não implantar, depois quem está sem agenda,
# depois quem só falta encerrar; quem já tem agenda marcada e os suspensos ficam por último
GRUPOS_PRIORIDADE = [
    "Crítico sem agenda",
    "Sem agenda marcada",
    "Encerrar",
    "Com agenda marcada",
    "Suspenso",
]


def _grupo_prioridade(r):
    if r.STATUS == "SUSPENSO":
        return "Suspenso"
    if pd.notna(r.PROXIMA_AGENDA):
        return "Com agenda marcada"
    if r.SAUDE == "Crítico":
        return "Crítico sem agenda"
    if r.SAUDE == "Encerrar":
        return "Encerrar"
    return "Sem agenda marcada"


def montar(caminho):
    bruto = ler_planilha(caminho)
    ref = bruto["atualizado_em"].normalize()
    hunter = _projetos_hunter(bruto["projetos"])

    # agendas de todos os projetos: a agenda do consultor inclui treinamento, consultoria, DBA etc.
    agenda_geral = _agendas(bruto["os"], bruto["planejada"], set(bruto["os"].COD_PROJETO), ref)
    info = bruto["projetos"].drop_duplicates("COD")[["COD", "SISTEMA", "TIPO_PROJETO"]]
    agenda_geral = agenda_geral.merge(info, left_on="COD_PROJETO", right_on="COD", how="left").drop(columns="COD")
    agenda_geral["SISTEMA_NOME"] = agenda_geral.SISTEMA.map(SISTEMAS).fillna("Outro")
    agenda_geral["HUNTER"] = agenda_geral.COD_PROJETO.isin(hunter.COD)
    agendas = agenda_geral[agenda_geral.HUNTER].drop(columns=["SISTEMA", "TIPO_PROJETO", "SISTEMA_NOME", "HUNTER"]).reset_index(drop=True)
    referencia = _referencia(hunter, agendas)

    horas = bruto["horas"][bruto["horas"].PROJETO.isin(hunter.COD)].copy()
    horas["GRUPO"] = horas.TIPO_HORA_ADQ.map(GRUPO_HORA).fillna("Contratadas")
    horas["QTDE"] = horas.QTDE_HORAS_ADQUIRIDAS.fillna(0)

    ab = hunter[hunter.STATUS.isin(STATUS_ABERTO)].copy().set_index("COD")
    ag = agendas[agendas.COD_PROJETO.isin(ab.index)]
    aconteceu = ag[ag.SITUACAO.isin(["Realizada", "Sem fechar"])]
    g = lambda df: df.groupby("COD_PROJETO")

    ab["ADQUIRIDAS"] = horas.groupby("PROJETO").QTDE.sum().reindex(ab.index).fillna(0)
    for grupo in ["Contratadas", "Bonificadas", "Transferidas", "Crédito financeiro"]:
        col = "HORAS_" + grupo.upper().replace(" ", "_").replace("É", "E")
        ab[col] = horas[horas.GRUPO == grupo].groupby("PROJETO").QTDE.sum().reindex(ab.index).fillna(0)
    ab["UTILIZADAS"] = g(ag[ag.SITUACAO == "Realizada"]).HORAS.sum().reindex(ab.index).fillna(0)
    ab["HORAS_SEM_FECHAR"] = g(ag[ag.SITUACAO == "Sem fechar"]).HORAS.sum().reindex(ab.index).fillna(0)
    ab["HORAS_AGENDADAS"] = g(ag[ag.SITUACAO == "Próxima"]).HORAS.sum().reindex(ab.index).fillna(0)
    ab["SALDO"] = ab.ADQUIRIDAS - ab.UTILIZADAS
    ab["SALDO_LIVRE"] = ab.SALDO - ab.HORAS_SEM_FECHAR - ab.HORAS_AGENDADAS
    ab["PCT_CONSUMO"] = (ab.UTILIZADAS / ab.ADQUIRIDAS).where(ab.ADQUIRIDAS > 0, 0)

    ab["AGENDAS_REALIZADAS"] = g(aconteceu).apply(lambda d: d[["OS", "DIA"]].drop_duplicates().shape[0]).reindex(ab.index).fillna(0)
    ab["PRIMEIRA_AGENDA"] = g(aconteceu).DIA.min().reindex(ab.index)
    ab["ULTIMA_AGENDA"] = g(aconteceu).DIA.max().reindex(ab.index)
    ab["PROXIMA_AGENDA"] = g(ag[ag.SITUACAO == "Próxima"]).DIA.min().reindex(ab.index)
    ab["AGENDAS_VENCIDAS"] = g(ag[ag.SITUACAO == "Vencida"]).size().reindex(ab.index).fillna(0)
    ab["AGENDAS_FUTURAS"] = g(ag[ag.SITUACAO == "Próxima"]).size().reindex(ab.index).fillna(0)
    ab["DIAS_EM_ABERTO"] = (ref - ab.DT_INICIAL).dt.days
    ab["DIAS_SEM_AGENDA"] = (ref - ab.ULTIMA_AGENDA.fillna(ab.DT_INICIAL)).dt.days
    ab["NUNCA_TEVE_AGENDA"] = ab.ULTIMA_AGENDA.isna()
    ab["FASE_ATUAL"] = g(aconteceu).FASE.max().reindex(ab.index).fillna(0).astype(int)
    ab["FASE_NOME"] = ab.FASE_ATUAL.map(FASES).fillna("Sem agenda de fase")
    ab["HORAS_PRESENCIAL"] = g(ag[(ag.SITUACAO == "Realizada") & (ag.EXECUCAO == "Presencial")]).HORAS.sum().reindex(ab.index).fillna(0)

    def consultores(d):
        s = d.groupby("CONSULTOR").HORAS.sum().sort_values(ascending=False)
        return [{"nome": n, "horas": round(float(h), 1)} for n, h in s.items()]
    ab["CONSULTORES"] = g(ag[ag.SITUACAO == "Realizada"]).apply(consultores).reindex(ab.index)
    ab["CONSULTORES"] = ab.CONSULTORES.apply(lambda v: v if isinstance(v, list) else [])

    # última OS executada (a agenda mais recente que aconteceu) e a próxima agenda marcada
    ultima = aconteceu.sort_values(["DIA", "OS"]).groupby("COD_PROJETO").tail(1).set_index("COD_PROJETO")
    proxima = ag[ag.SITUACAO == "Próxima"].sort_values(["DIA", "OS"]).groupby("COD_PROJETO").head(1).set_index("COD_PROJETO")
    for prefixo, df in (("ULTIMA_OS", ultima), ("PROXIMA_OS", proxima)):
        ab[prefixo] = df.OS.reindex(ab.index)
        ab[prefixo + "_ETAPA"] = df.ETAPA.reindex(ab.index)
        ab[prefixo + "_CONSULTOR"] = df.CONSULTOR.reindex(ab.index)
        ab[prefixo + "_HORAS"] = df.HORAS.reindex(ab.index)
        ab[prefixo + "_EXECUCAO"] = df.EXECUCAO.reindex(ab.index)
    ab["ULTIMA_OS_SITUACAO"] = ultima.SITUACAO.reindex(ab.index)

    p75 = referencia.dias_p75.to_dict()
    saude = ab.apply(lambda r: _saude(r, p75), axis=1)
    ab["SAUDE"] = saude.str[0]
    ab["MOTIVOS"] = saude.str[1]

    ab["PRIORIDADE_GRUPO"] = ab.apply(_grupo_prioridade, axis=1)
    gravidade = ab.SAUDE.map({s: i for i, s in enumerate(SAUDES)}).fillna(-1)
    ordem = (ab.assign(_g=ab.PRIORIDADE_GRUPO.map(GRUPOS_PRIORIDADE.index), _s=-gravidade, _d=-ab.DIAS_SEM_AGENDA)
               .sort_values(["_g", "_s", "_d"]).index)
    ab.loc[ordem, "PRIORIDADE"] = range(1, len(ordem) + 1)
    ab["PRIORIDADE"] = ab.PRIORIDADE.astype(int)

    return Modelo(ref=ref, atualizado_em=bruto["atualizado_em"], hunter=hunter,
                  abertos=ab.reset_index(), agendas=agendas, horas=horas, referencia=referencia,
                  agenda_geral=agenda_geral, fases_referencia=_fases_referencia(hunter, agendas))
